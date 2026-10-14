"""Resume text extraction for PDF, DOCX and TXT inputs.

We never trust the client-supplied filename or MIME type on its own: the
extension decides which parser runs, but the actual byte content is inspected
first (magic numbers, decodability, structure) and rejected when it does not
match.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pdfplumber
from docx import Document  # python-docx

from ..config import get_settings
from ..core.errors import ApiError

ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt")

_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK\x03\x04"

# ``(cid:NNN)`` sequences are produced by pdfminer when a PDF font has no
# ToUnicode mapping for a glyph. Bullets and typographic punctuation are the most
# common occurrences; the remainder are stripped so they cannot hide content.
_CID_RE = re.compile(r"\(cid:\s*(\d+)\)")
_CID_GLYPHS: dict[int, str] = {
    32: " ",  # space
    127: "\u2022",  # bullet (WinAnsi / dingbats)
    149: "\u2022",
    183: "\u00b7",
    150: "\u2013",  # en dash
    151: "\u2014",  # em dash
    145: "\u2018",
    146: "\u2019",
    147: "\u201c",
    148: "\u201d",
}

# Bytes that should never appear in a plain-text resume.
_TEXT_DISALLOWED_CONTROL_RATIO = 0.30

# A word broken across two PDF text lines with a trailing hyphen. Joined only
# when there is clear evidence of line-wrap hyphenation: at least four letters
# before the hyphen and a short (2-4 letter) lowercase continuation. Genuine
# hyphens that fall mid-line are never touched. A genuine compound that happens
# to be split exactly at its hyphen by the layout engine is the known ambiguous
# case; requiring a short continuation keeps false joins rare.
_HYPHEN_WRAP_RE = re.compile(r"([A-Za-z]{4,})-\n([a-z]{2,4})\b")
# A second pattern for camel-case tokens split at a letter boundary, e.g.
# "Post-\ngreSQL" -> "PostgreSQL". The continuation starts lowercase but
# contains an interior capital, which ordinary word-wrap hyphenation never does.
_HYPHEN_WRAP_CAMEL_RE = re.compile(r"([A-Za-z]{2,})-\n([a-z][A-Za-z]*[A-Z][A-Za-z]*)\b")

# f-ligatures produced by TeX/Word fonts; normalised to ASCII so that words like
# "classi<fi>cation" match a plain-text search and the ATS checks.
_LIGATURES: dict[str, str] = {
    "\ufb00": "ff",
    "\ufb01": "fi",
    "\ufb02": "fl",
    "\ufb03": "ffi",
    "\ufb04": "ffl",
    "\ufb05": "st",
    "\ufb06": "st",
}


def _join_hyphenated_wraps(text: str) -> str:
    """Re-join words that a PDF text layer split with a trailing hyphen.

    Handles constructs such as ``"... scalabil-\\nity ..."`` -> ``"scalability"``
    and camel-case tokens such as ``"Post-\\ngreSQL"`` -> ``"PostgreSQL"``.
    Only line-wrap hyphenation is joined; hyphens inside a line and genuine
    hyphenated compounds are preserved.
    """
    text = _HYPHEN_WRAP_RE.sub(lambda match: match.group(1) + match.group(2), text)
    return _HYPHEN_WRAP_CAMEL_RE.sub(lambda match: match.group(1) + match.group(2), text)


@dataclass
class ExtractionResult:
    """The outcome of parsing a single uploaded file."""

    text: str
    parser: str
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)


def _extension_of(filename: str | None) -> str:
    if not filename:
        raise ApiError(415, "unsupported_file_type", "A filename with an extension is required.")
    return Path(filename).suffix.lower()


def _looks_like_text(data: bytes) -> bool:
    if not data:
        return False
    control = sum(1 for b in data if b < 0x09 or (0x0E <= b <= 0x1F))
    return (control / len(data)) <= _TEXT_DISALLOWED_CONTROL_RATIO


def _decode_text(data: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding), warnings
        except UnicodeDecodeError:
            if encoding == "utf-8":
                warnings.append("File was not valid UTF-8; falling back to a legacy encoding.")
            continue
    raise ApiError(422, "invalid_content", "The text file could not be decoded.")


def _normalize_text(text: str) -> str:
    """Clean common PDF/DOCX text-layer artefacts from extracted text.

    * ``(cid:NNN)`` glyph escapes are mapped to printable punctuation where the
      codepoint is known (bullet, en/em dash) and dropped otherwise.
    * Zero-width and soft formatting marks are removed.
    * Non-breaking spaces become ordinary spaces so word splitting works.
    * Ligatures (``ﬁ``, ``ﬂ`` …) are expanded to plain letters.
    """
    text = text.replace("\ufeff", "").replace("\r\n", "\n").replace("\r", "\n")
    for char in ("\u200b", "\u200c", "\u200d", "\u200e", "\u200f", "\u202a", "\u202c"):
        text = text.replace(char, "")
    for char in ("\u00a0", "\u202f", "\u2009", "\u2007"):
        text = text.replace(char, " ")
    for ligature, replacement in _LIGATURES.items():
        text = text.replace(ligature, replacement)

    def _replace_cid(match: re.Match[str]) -> str:
        glyph = _CID_GLYPHS.get(int(match.group(1)))
        return glyph if glyph is not None else ""

    return _CID_RE.sub(_replace_cid, text)


def _completion_warnings(text: str, page_count: int | None) -> list[str]:
    """Surface signals that text extraction may have lost content."""
    warnings: list[str] = []
    words = len(text.split())
    if page_count and page_count > 0:
        per_page = words / page_count
        if per_page < 35 and words < 200:
            warnings.append(
                f"Text extraction looks incomplete: only about {words} words were "
                f"recovered from {page_count} page(s). The PDF may use fonts or a "
                "layout that are hard to read programmatically. Consider pasting the "
                "resume text instead."
            )
    suspicious_space_merges = 0
    for line in text.splitlines():
        stripped = line.strip()
        letters = sum(1 for char in stripped if char.isalpha())
        if len(stripped) >= 50 and " " not in stripped and letters >= 12:
            suspicious_space_merges += 1
    if suspicious_space_merges >= 3:
        warnings.append(
            "Some extracted lines appear to be missing the spaces between words. "
            "This can happen when a PDF does not encode spaces normally; pasting the "
            "resume text will preserve it exactly."
        )
    return warnings


def _extract_txt(data: bytes) -> ExtractionResult:
    if not _looks_like_text(data):
        raise ApiError(
            415,
            "unsupported_file_type",
            "The .txt file does not appear to contain readable text.",
        )
    text, warnings = _decode_text(data)
    return ExtractionResult(text=text, parser="plain-text", warnings=warnings)


def _extract_docx(data: bytes) -> ExtractionResult:
    if not data.startswith(_ZIP_MAGIC):
        raise ApiError(415, "invalid_content", "The .docx file is not a valid Office document.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ApiError(
                    415, "invalid_content", "The .docx file is not a valid Word document."
                )
        document = Document(io.BytesIO(data))
    except ApiError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ApiError(415, "invalid_content", "The .docx file could not be read.") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = _normalize_text("\n".join(parts))
    return ExtractionResult(text=text, parser="python-docx")


def _extract_pdf(data: bytes) -> ExtractionResult:
    if not data.startswith(_PDF_MAGIC):
        raise ApiError(415, "invalid_content", "The .pdf file is not a valid PDF document.")
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(_normalize_text(page.extract_text() or "")) for page in pdf.pages]
            page_count = len(pdf.pages)
    except ApiError:
        raise
    except Exception as exc:  # pdfminer raises a variety of errors on corrupt input
        raise ApiError(415, "invalid_content", "The .pdf file could not be read.") from exc

    text = "\n".join(pages)
    text = _join_hyphenated_wraps(text)
    if not text.strip():
        # A structurally valid PDF with no text layer is almost certainly a scan.
        raise ApiError(
            422,
            "ocr_not_supported",
            "No selectable text was found in this PDF. It may be a scanned image; "
            "OCR is not supported yet. Please upload a text-based PDF or paste the resume text.",
        )
    warnings = [] if page_count else ["The PDF reported zero pages."]
    warnings.extend(_completion_warnings(text, page_count))
    return ExtractionResult(
        text=text, parser="pdfplumber", page_count=page_count, warnings=warnings
    )


_EXTRACTORS = {
    ".txt": _extract_txt,
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
}


def extract_text(filename: str | None, data: bytes) -> ExtractionResult:
    """Validate and extract text from an uploaded file.

    Raises ``ApiError`` for unsupported types, oversized/invalid content, or
    files from which no text could be recovered.
    """
    settings = get_settings()
    extension = _extension_of(filename)

    if extension not in ALLOWED_EXTENSIONS:
        raise ApiError(
            415,
            "unsupported_file_type",
            "Unsupported file type. Upload a PDF, DOCX or TXT file.",
            [{"allowed_extensions": list(ALLOWED_EXTENSIONS)}],
        )

    if len(data) == 0:
        raise ApiError(400, "empty_resume", "The uploaded file is empty.")

    if len(data) > settings.max_upload_bytes:
        max_mb = settings.max_upload_bytes // (1024 * 1024)
        raise ApiError(
            413,
            "file_too_large",
            f"File is too large. The maximum size is {max_mb} MB.",
        )

    result = _EXTRACTORS[extension](data)

    if len(result.text.strip()) < 20:
        raise ApiError(
            422,
            "empty_resume",
            "No usable resume text could be extracted from the file.",
        )
    return result
